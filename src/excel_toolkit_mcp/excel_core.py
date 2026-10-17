#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
excel_toolkit.py — 纯本地 Excel 工具包(无 API Key、无 LLM、不联网)。

两大功能:
  1) 公式生成  formula:用自然语言描述需求,返回可用的 Excel/Google Sheets 公式。
     基于"意图模式 + 参数槽位"的确定性引擎(不是调用大模型),覆盖最高频的表格计算,
     结果可直接粘贴进单元格,并给中文解释。
  2) 表格清洗  clean:读 CSV/XLSX,做一组可组合的清洗动作,输出干净的 .xlsx/.csv。
     动作:去重/去首尾空格/统一日期/统一金额数字/空值填充/列名规范化/选择列/筛选/拆分列。

为什么有用:不把财务表传给在线 AI,本地处理敏感表格;高频需求零等待、结果确定可复核。

用法:
  # 公式:交互式一句话
  python3 excel_toolkit.py formula "计算B2到B100的平均值,忽略空单元格"
  python3 excel_toolkit.py formula --list                    # 列出支持的意图
  # 清洗:用一个清洗指令 JSON
  python3 excel_toolkit.py clean messy.xlsx clean.json -o clean.xlsx
  python3 excel_toolkit.py clean data.csv clean.json --csv -o out.csv
  # 自检
  python3 excel_toolkit.py --selftest
依赖:openpyxl (必需,处理 xlsx);csv 为标准库。
"""
import argparse, csv, json, os, re, sys, datetime

try:
    import openpyxl
except ImportError:
    print("缺少 openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(2)


# ---------------- 公式生成引擎 ----------------
# 每条:(意图关键词正则, 参数抽取函数, 公式模板, 说明)
def _rng(g, default=""):
    m = re.search(r"([A-Z]{1,3}\d+)\s*(?:到|至|-|–|:|—)\s*([A-Z]{1,3}\d+)", g)
    if m:
        return f"{m.group(1)}:{m.group(2)}"
    m2 = re.search(r"([A-Z]{1,3})\s*(?:列)?\s*(?:的|从)?\s*(?:第)?\s*(\d+)?\s*(?:行)?\s*(?:到|至|-|–|:|—)\s*(?:第)?\s*(\d+)?\s*行?", g)
    if m2 and m2.group(2) and m2.group(3):
        return f"{m2.group(1)}{m2.group(2)}:{m2.group(1)}{m2.group(3)}"
    if m2 and m2.group(1):
        return f"{m2.group(1)}:{m2.group(1)}"
    return default


def _cell(g):
    m = re.search(r"\b([A-Z]{1,3}\d+)\b", g)
    return m.group(1) if m else "A1"


def _col(g):
    m = re.search(r"\b([A-Z]{1,3})\s*列?\b", g)
    if m:
        return m.group(1)
    m = re.search(r"([A-Z]{1,3})\d*", g)
    return m.group(1) if m else "A"


INTENTS = [
    # 顺序敏感:复合/条件意图必须在简单"计数/取字符"之前,避免"条件计数"被"个/计数"抢匹配
    (r"多条件求和|sumifs", lambda g: ("=SUMIFS(求和列2:求和列100,条件列2:条件列100,\"条件1\")"),
     "多条件求和模板,按实际列替换"),
    (r"条件求和|满足.*求和|sumif", lambda g: (f"=SUMIF({_col(g)}2:{_col(g)}100,\"条件\",求和列2:求和列100)"),
     "按单条件求和;请把\"条件\"换成如 \">1000\" 或 \"已完成\",并指定求和列"),
    (r"条件计数|满足.*几个|满足.*多少|countif|大于.*几个|大于.*多少", lambda g: (f"=COUNTIF({_col(g)}2:{_col(g)}100,\"条件\")"),
     "按条件计数;如 =COUNTIF(A2:A100,\">=60\")"),
    (r"非空.*(计数|个数)|多少条记录", lambda g: (f"=COUNTA({_rng(g,'B2:B100')})"), "非空单元格计数"),
    (r"百分比排名|占比", lambda g: (f"={_cell(g)}/SUM({_rng(g,'B$2:B$100')})"),
     "单项占区域合计的比例(把单元格格式设为百分比)"),
    (r"if.*判断|如果.*那么|条件判断|大于.*显示", lambda g: ("=IF(A1>=60,\"达标\",\"未达标\")"),
     "条件判断模板;把 A1>=60 与两个结果替换成你的规则"),
    (r"vlookup|查表|查第?\s*\d+\s*列|按.*查找|匹配.*返回|在.*表.*查", lambda g: ("=VLOOKUP(查找值, 表区域, 返回第几列, FALSE)"),
     "精确查表;FALSE 表示精确匹配。例 =VLOOKUP(D2,A:C,3,FALSE)"),
    (r"xlookup", lambda g: ("=XLOOKUP(查找值, 查找列, 返回列)"), "新式查表(Excel 365)"),
    (r"拼接.*文本|合并.*文本|拼起来|连起来|concat|join", lambda g: (f"=TEXTJOIN(\" \",TRUE,{_rng(g,'A2:A10')})"),
     "用分隔符拼接非空文本(此处为空格,可改 \",\")"),
    (r"同比|环比增长|增长率", lambda g: ("=(本期-上期)/上期"), "增长率模板,替换本期/上期单元格"),
    (r"日期差|相差多少天|天数差|datedif", lambda g: ("=DATEDIF(开始日期,结束日期,\"d\")"),
     "两日期间隔天数(\"m\"月 / \"y\"年)"),
    (r"求和|总和|加总|合计|sum", lambda g: (f"=SUM({_rng(g,'A1:A100')})"),
     "对区域求和"),
    (r"平均值|均值|平均|average|avg", lambda g: (f"=AVERAGE({_rng(g,'B2:B100')})"),
     "平均值,自动忽略文本与空单元格"),
    (r"中位数|median", lambda g: (f"=MEDIAN({_rng(g,'B2:B100')})"), "中位数"),
    (r"众数|出现最多", lambda g: (f"=MODE({_rng(g,'B2:B100')})"), "众数(出现最多的值)"),
    (r"最大值|最大|max", lambda g: (f"=MAX({_rng(g,'B2:B100')})"), "最大值"),
    (r"最小值|最小|min", lambda g: (f"=MIN({_rng(g,'B2:B100')})"), "最小值"),
    (r"计数|多少个|个数|count", lambda g: (f"=COUNT({_rng(g,'B2:B100')})"),
     "计数(只数数字);数非空用 COUNTA,数空白用 COUNTBLANK"),
    (r"标准差|波动|standard dev", lambda g: (f"=STDEV({_rng(g,'B2:B100')})"), "样本标准差"),
    (r"方差|variance", lambda g: (f"=VAR({_rng(g,'B2:B100')})"), "样本方差"),
    (r"去重|不重复值|unique", lambda g: (f"=UNIQUE({_rng(g,'A2:A100')})"),
     "返回区域中的不重复值(Excel 365/Google Sheets 动态数组)"),
    (r"排名|rank", lambda g: (f"=RANK({_cell(g)},{_rng(g,'B$2:B$100')},0)"),
     "某值在区域中的降序排名(第三参 1 为升序)"),
    (r"取左边|左.*字符|left", lambda g: (f"=LEFT({_cell(g)},3)"), "从左侧取 N 个字符(示例 3)"),
    (r"取右边|右.*字符|right", lambda g: (f"=RIGHT({_cell(g)},3)"), "从右侧取 N 个字符(示例 3)"),
    (r"长度|几个字|len", lambda g: (f"=LEN({_cell(g)})"), "文本字符长度"),
    (r"去空格|trim", lambda g: (f"=TRIM({_cell(g)})"), "去除多余空格(首尾+中间连续)"),
    (r"今天|当前日期|today", lambda g: ("=TODAY()"), "当前日期(随系统更新)"),
    (r"现在.*时间|当前时间|now", lambda g: ("=NOW()"), "当前日期时间"),
    (r"四舍五入|round", lambda g: (f"=ROUND({_cell(g)},2)"), "四舍五入到 N 位小数(示例 2)"),
    (r"取整|int", lambda g: (f"=INT({_cell(g)})"), "向下取整"),
]


def formula(text):
    t = text.lower().replace("（", "(").replace("）", ")")
    for pat, fn, desc in INTENTS:
        if re.search(pat, t):
            return {"formula": fn(text), "intent": desc}
    return {"formula": None,
            "hint": "未识别意图。用 --list 看支持项,或把需求说得更具体(指明区域,如\"B2到B100的平均值\")。"}


# ---------------- 表格清洗引擎 ----------------
def _norm_header(h):
    h = str(h).strip()
    h = re.sub(r"\s+", "_", h)
    h = re.sub(r"[（(].*?[)）]", "", h)
    return h.strip("_") or "col"


_DATE_PATTERNS = [
    ("%Y-%m-%d", re.compile(r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}$")),
    ("%d/%m/%Y", re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")),
    ("%Y年%m月%d日", re.compile(r"^\d{4}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日?$")),
]


def _to_iso_date(v):
    if v is None:
        return v
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v).strip()
    # 归一化分隔符:2026.09.03 / 2026/9/1 -> 2026-09-03
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})$", s)
    if m:
        y, mo, da = (int(x) for x in m.groups())
        try:
            return datetime.date(y, mo, da).isoformat()
        except ValueError:
            return s
    m = re.match(r"^(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日?$", s)
    if m:
        y, mo, da = (int(x) for x in m.groups())
        try:
            return datetime.date(y, mo, da).isoformat()
        except ValueError:
            return s
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", s)
    if m:
        da, mo, y = (int(x) for x in m.groups())
        try:
            return datetime.date(y, mo, da).isoformat()
        except ValueError:
            return s
    return s


def _to_number(v):
    if v is None:
        return v
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", "").replace("，", "")
    s = s.replace("￥", "").replace("¥", "").replace("$", "").replace("元", "").replace("%", "")
    s = re.sub(r"\s+", "", s)
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        f = float(s)
        return int(f) if f.is_integer() else round(f, 6)
    return v


def clean_rows(rows, spec):
    if not rows:
        return rows, []
    headers = [str(h) for h in rows[0]]
    data = [list(r) for r in rows[1:]]

    def ridx(name):
        name = name.strip()
        if name in headers:
            return headers.index(name)
        nh = _norm_header(name)
        for i, h in enumerate(headers):
            if _norm_header(h) == nh:
                return i
        return None

    if spec.get("normalize_headers", True):
        headers = [_norm_header(h) for h in headers]

    # 去空列(整列无值)
    if spec.get("drop_empty_cols", True):
        keep = [j for j in range(len(headers))
                if any(str(r[j]).strip() != "" for r in data if j < len(r))]
        headers = [headers[j] for j in keep]
        data = [[r[j] if j < len(r) else "" for j in keep] for r in data]

    # 逐单元格
    trim = spec.get("strip_whitespace", True)
    date_cols = set(i for c in spec.get("date_columns", []) if (i := ridx(c)) is not None)
    num_cols = set(i for c in spec.get("number_columns", []) if (i := ridx(c)) is not None)
    fill = spec.get("fillna")  # {列: 值} 或 {"*": "未知"}
    notes = []
    for r in data:
        for j in range(len(r)):
            v = r[j]
            if trim and isinstance(v, str):
                v = v.strip()
            if j in date_cols:
                v = _to_iso_date(v)
            if j in num_cols:
                v = _to_number(v)
            if fill is not None and (v is None or str(v).strip() == ""):
                key = headers[j] if j < len(headers) else "*"
                v = fill.get(key, fill.get("*", v))
            r[j] = v

    # 拆分列:在逐单元格处理之后,统一拆,并把每行补齐到新表头长度
    for sp in spec.get("split_columns", []):
        si = ridx(sp["column"])
        if si is None:
            continue
        sep = sp.get("separator", r"[,\s]+")
        into = sp.get("into", [f"{headers[si]}_{i+1}" for i in range(sp.get("max", 3))])
        # 追加尚不存在的新列
        new_cols = [n for n in into if n not in headers]
        base_len = len(headers)
        headers = headers + new_cols
        for r in data:
            while len(r) < base_len:
                r.append("")
            src = str(r[si] if si < len(r) else "")
            parts = re.split(sep, src.strip())
            # 先给新列补位
            r.extend([""] * len(new_cols))
            for k, name in enumerate(into):
                val = parts[k] if k < len(parts) else ""
                r[headers.index(name)] = val

    # 选择列（在拆分列之后执行：这样才能选到 split_columns 新生成的市/省等列）
    if spec.get("select_columns"):
        idxs = [ridx(c) for c in spec["select_columns"]]
        idxs = [i for i in idxs if i is not None]
        headers = [headers[i] for i in idxs]
        data = [[r[i] if i < len(r) else "" for i in idxs] for r in data]

    # 筛选(简单等值/比较)
    flt = spec.get("filter")
    if flt:
        fi = ridx(flt["column"])
        op = flt.get("op", "eq")
        fval = flt.get("value")
        def ok(x):
            xn, vn = _to_number(x), _to_number(fval)
            num = isinstance(xn, (int, float)) and isinstance(vn, (int, float))
            try:
                if op == "contains":
                    return str(vn) in str(xn)
                if op in ("eq", "ne"):
                    same = (xn == vn) if num else (str(xn).strip() == str(vn).strip())
                    return same if op == "eq" else not same
                # 比较运算仅对数值生效，文本无法比较时保守返回 False
                if not num:
                    return False
                return {"gt": xn > vn, "lt": xn < vn, "ge": xn >= vn, "le": xn <= vn}[op]
            except Exception:
                return False
        before = len(data)
        data = [r for r in data if fi is not None and fi < len(r) and ok(r[fi])]
        notes.append(f"filter {flt.get('column')} {op} {fval}: {before}→{len(data)} 行")

    # 去重
    if spec.get("drop_duplicates", True):
        seen, uniq = set(), []
        key_cols = [ridx(c) for c in spec.get("dedup_subset", [])]
        key_cols = [i for i in key_cols if i is not None]
        for r in data:
            k = tuple(r[i] if i < len(r) else "" for i in key_cols) if key_cols else tuple(r)
            if k not in seen:
                seen.add(k); uniq.append(r)
        if len(uniq) != len(data):
            notes.append(f"去重 {len(data)}→{len(uniq)} 行")
        data = uniq

    return [headers] + data, notes


def read_table(path):
    if path.lower().endswith(".csv"):
        # 尝试常见编码
        for enc in ("utf-8-sig", "gbk", "latin-1"):
            try:
                with open(path, encoding=enc, newline="") as f:
                    return [list(r) for r in csv.reader(f)]
            except UnicodeDecodeError:
                continue
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def write_table(rows, out, as_csv):
    if as_csv or out.lower().endswith(".csv"):
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
    else:
        wb = openpyxl.Workbook(); ws = wb.active
        for r in rows:
            ws.append(["" if c is None else c for c in r])
        wb.save(out)


def _selftest():
    # 公式
    assert formula("计算B2到B100的平均值，忽略空单元格")["formula"] == "=AVERAGE(B2:B100)"
    assert "SUM" in formula("A列从第2行到第50行求和")["formula"]
    assert formula("vlookup按D2在A到B表查第2列")["formula"].startswith("=VLOOKUP")
    assert formula("统计A列满足大于60的有几个")["formula"].startswith("=COUNTIF"), formula("统计A列满足大于60的有几个")
    assert formula("按A列对B列满足已完成的求和")["formula"].startswith("=SUMIF")
    assert formula("把A2到A5的文本用空格拼起来")["formula"].startswith("=TEXTJOIN")
    assert formula("随便聊聊天气")["formula"] is None
    # 日期归一化(点分隔/中文/斜杠)
    assert _to_iso_date("2026.09.03") == "2026-09-03"
    assert _to_iso_date("2026年9月1日") == "2026-09-01"
    assert _to_iso_date("2026/9/1") == "2026-09-01"
    # 清洗
    rows = [[" 姓名（Name） ", "金额", "日期", "地区"],
            [" 张三 ", "￥1,200.50", "2026年9月1日", "郑州,河南"],
            ["张三", "1200.5", "2026-09-01", "郑州,河南"],
            ["李四", "$99", "2026.09.02", "北京,北京"]]
    spec = {"number_columns": ["金额"], "date_columns": ["日期"],
            "split_columns": [{"column": "地区", "into": ["市", "省"]}]}
    out, notes = clean_rows(rows, spec)
    head = out[0]
    assert "姓名" == head[0], head
    assert out[1][1] == 1200.5 and out[1][2] == "2026-09-01", out[1]
    assert out[2][2] == "2026-09-02", out[2]
    assert len(out) == 3, f"去重后应 3 行(表头+张三+李四): {out}"
    assert "市" in head and "省" in head, head
    assert out[1][head.index("市")] == "郑州" and out[1][head.index("省")] == "河南", out[1]
    assert out[2][head.index("市")] == "北京" and out[2][head.index("省")] == "北京", out[2]
    print("selftest PASS: 公式(顺序/条件/未识别) + 日期三种格式 + 清洗(数字/去重/拆列对齐/表头) 全通过")
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("command", nargs="?", choices=["formula", "clean"])
    ap.add_argument("args", nargs="*")
    ap.add_argument("-o", "--output")
    ap.add_argument("--csv", action="store_true")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()

    if a.selftest:
        sys.exit(_selftest())
    if a.command == "formula":
        if a.list:
            for _, _, desc in INTENTS:
                print(" -", desc)
            return
        if not a.args:
            ap.error("formula 需要一句需求描述")
        print(json.dumps(formula(" ".join(a.args)), ensure_ascii=False, indent=2))
    elif a.command == "clean":
        if len(a.args) < 2:
            ap.error("clean 需要: 输入文件 清洗指令.json [-o 输出]")
        inp, specp = a.args[0], a.args[1]
        spec = json.load(open(specp, encoding="utf-8"))
        rows = read_table(inp)
        out, notes = clean_rows(rows, spec)
        outp = a.output or ("cleaned.xlsx" if not a.csv else "cleaned.csv")
        write_table(out, outp, a.csv)
        print("OUT:", os.path.abspath(outp), "| 行数(含表头):", len(out))
        for n in notes:
            print(" -", n)
    else:
        ap.print_help()


if __name__ == "__main__":
    main()
